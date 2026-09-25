from __future__ import annotations

from dataclasses import dataclass
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.fonts import addMapping
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.models.analysis import Analysis
from app.services.new_summaries import with_summary_display_stage


PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXPORT_FORMATS = {"pdf", "docx"}

_ACCENT = "#0B6F54"
_SUCCESS = "#0FA36B"
_WARNING = "#C77800"
_DANGER = "#D92D20"
_MUTED = "#5D6675"
_LINE = "#DDE3EA"
_SURFACE = "#F6F8FA"
_TEXT = "#111827"


class NewSummaryExportUnavailableError(Exception):
    pass


class UnsupportedNewSummaryExportFormatError(Exception):
    pass


@dataclass(frozen=True)
class NewSummaryExport:
    content: bytes
    filename: str
    media_type: str

    @property
    def content_disposition(self) -> str:
        quoted = quote(self.filename)
        return f"attachment; filename*=UTF-8''{quoted}"


@dataclass(frozen=True)
class NewSummaryExportProvenance:
    analysis_id: str
    document_id: str
    skill_name: str
    skill_version: str
    provider: str
    model: str
    source_revision: str | None


def build_new_summary_export(*, analysis: Analysis, file_format: str, display_stage: str | None = None) -> NewSummaryExport:
    normalized_format = file_format.lower()
    if normalized_format not in EXPORT_FORMATS:
        raise UnsupportedNewSummaryExportFormatError(file_format)

    report = _read_completed_report(analysis)
    if display_stage is not None:
        report = {
            **report,
            "ru": with_summary_display_stage(report["ru"], display_stage),
            "en": with_summary_display_stage(report["en"], display_stage),
        }
    provenance = _provenance(analysis=analysis, source_revision=report["source_revision"])
    title = _clean_text(report["ru"].get("title") or report["en"].get("title") or "AI Summary")
    filename = f"{_safe_filename(title)}.{normalized_format}"
    if normalized_format == "docx":
        return NewSummaryExport(
            content=_build_docx(report, provenance),
            filename=filename,
            media_type=DOCX_MEDIA_TYPE,
        )
    return NewSummaryExport(
        content=_build_pdf(report, provenance),
        filename=filename,
        media_type=PDF_MEDIA_TYPE,
    )


def _read_completed_report(analysis: Analysis) -> dict[str, Any]:
    output = analysis.structured_output if isinstance(analysis.structured_output, dict) else {}
    result = output.get("result") if isinstance(output.get("result"), dict) else {}
    state = result.get("new_summary") if isinstance(result.get("new_summary"), dict) else {}
    versions: dict[str, Any] = {"source_revision": state.get("source_revision") if isinstance(state.get("source_revision"), str) else None}
    for language in ("ru", "en"):
        variant = state.get(language) if isinstance(state.get(language), dict) else {}
        payload = variant.get("payload") if isinstance(variant.get("payload"), dict) else None
        if variant.get("status") != "completed" or payload is None:
            raise NewSummaryExportUnavailableError("New Summary is not completed")
        versions[language] = payload
    return versions


def _provenance(*, analysis: Analysis, source_revision: str | None) -> NewSummaryExportProvenance:
    return NewSummaryExportProvenance(
        analysis_id=str(analysis.id),
        document_id=str(analysis.document_id),
        skill_name=str(analysis.skill_id),
        skill_version=str(analysis.skill_version),
        provider=str(analysis.provider),
        model=str(analysis.model),
        source_revision=source_revision,
    )


def _build_docx(report: dict[str, Any], provenance: NewSummaryExportProvenance) -> bytes:
    document = DocxDocument()
    document.core_properties.title = _clean_text(report["ru"].get("title") or report["en"].get("title") or "AI Summary")
    document.core_properties.subject = "Gate Challenger AI Summary export"
    document.core_properties.comments = _short_provenance_text(provenance)
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    for index, language in enumerate(("en", "ru")):
        if index:
            document.add_page_break()
        _append_docx_version(document, report[language], language)
    _append_docx_provenance(document, provenance)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _append_docx_version(document: DocxDocument, content: dict[str, Any], language: str) -> None:
    labels = _labels(language)
    title = _clean_text(content.get("title") or "AI Summary")

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(17, 24, 39)

    stage = document.add_paragraph()
    stage_run = stage.add_run(f"{labels['stage']}: ")
    stage_run.bold = True
    stage.add_run(_clean_text(content.get("stage") or "Unknown"))

    _append_docx_traction(document, content, labels)
    _append_docx_text_section(document, labels["context"], [_clean_text(content.get("context"))])
    _append_docx_required(document, content, labels)
    _append_docx_list_section(document, labels["critical"], _string_list(content.get("critical_problems")), RGBColor(185, 28, 28))
    _append_docx_list_section(document, labels["other"], _string_list(content.get("other")), RGBColor(93, 102, 117))
    _append_docx_details(document, content, labels)


def _append_docx_heading(document: DocxDocument, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.space_before = Pt(12)
    paragraph.space_after = Pt(5)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(17, 24, 39)


def _append_docx_text_section(document: DocxDocument, title: str, items: list[str]) -> None:
    values = [item for item in items if item]
    if not values:
        return
    _append_docx_heading(document, title)
    for item in values:
        paragraph = document.add_paragraph(item)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15


def _append_docx_list_section(document: DocxDocument, title: str, items: list[str], color: RGBColor) -> None:
    if not items:
        return
    _append_docx_heading(document, title)
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.color.rgb = color


def _append_docx_traction(document: DocxDocument, content: dict[str, Any], labels: dict[str, str]) -> None:
    traction = content.get("traction_summary")
    if not _has_traction(traction):
        return
    _append_docx_heading(document, labels["traction"])
    periods = [str(item) for item in traction["periods"]]
    rows = traction["rows"]
    table = document.add_table(rows=1 + len(rows), cols=1 + len(periods))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = _clean_text(traction.get("metric_label"))
    for index, period in enumerate(periods, start=1):
        header[index].text = period
    for row_index, row in enumerate(rows, start=1):
        cells = table.rows[row_index].cells
        cells[0].text = _clean_text(row.get("label"))
        values = row.get("values") if isinstance(row.get("values"), list) else []
        for cell_index, _period in enumerate(periods, start=1):
            cells[cell_index].text = _clean_text(values[cell_index - 1] if cell_index - 1 < len(values) else "")
    _style_docx_table(table)


def _style_docx_table(table) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True


def _append_docx_required(document: DocxDocument, content: dict[str, Any], labels: dict[str, str]) -> None:
    items = content.get("required_elements")
    if not isinstance(items, list) or not items:
        return
    _append_docx_heading(document, labels["required"])
    for item in items:
        if not isinstance(item, dict):
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.space_after = Pt(2)
        title = paragraph.add_run(_clean_text(item.get("label")))
        title.bold = True
        status = _localized_status(item, labels)
        status_run = paragraph.add_run(f" — {status}")
        status_run.bold = True
        status_run.font.color.rgb = {
            "present": RGBColor(15, 163, 107),
            "partial": RGBColor(199, 120, 0),
            "missing": RGBColor(217, 45, 32),
            "fraction": RGBColor(17, 24, 39),
        }[_required_status_kind(item)]
        evidence = document.add_paragraph(_clean_text(item.get("evidence")))
        evidence.paragraph_format.left_indent = Inches(0.22)
        evidence.paragraph_format.space_after = Pt(5)
        for run in evidence.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(93, 102, 117)


def _append_docx_details(document: DocxDocument, content: dict[str, Any], labels: dict[str, str]) -> None:
    entries = _ordered_details(content)
    if not entries:
        return
    _append_docx_heading(document, labels["appendices"])
    for index, (detail_id, detail) in enumerate(entries, start=1):
        _append_docx_heading(document, _appendix_title(content, detail_id, index))
        _append_docx_detail(document, detail, labels)


def _append_docx_detail(document: DocxDocument, detail: dict[str, Any], labels: dict[str, str]) -> None:
    detail_type = detail.get("type")
    if detail_type == "solution_validation":
        for item in _dict_list(detail.get("items")):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(_clean_text(item.get("text")))
            status_run = paragraph.add_run(f" — {_verdict_status(item.get('verdict'), labels)}")
            status_run.bold = True
            status_run.font.color.rgb = RGBColor(15, 163, 107) if item.get("verdict") == "confirmed" else RGBColor(199, 120, 0)
        return
    if detail_type == "metric_binding":
        for title, key in ((labels["input_metrics"], "input_metrics"), (labels["output_metrics"], "output_metrics")):
            metrics = _dict_list(detail.get(key))
            if not metrics:
                continue
            paragraph = document.add_paragraph()
            run = paragraph.add_run(title)
            run.bold = True
            for item in metrics:
                bullet = document.add_paragraph(style="List Bullet")
                bullet.add_run(_clean_text(item.get("metric"))).bold = True
                if item.get("evidence"):
                    bullet.add_run(f" - {_clean_text(item.get('evidence'))}")
                status_run = bullet.add_run(f" — {_detail_status(item.get('binding'), labels)}")
                status_run.bold = True
                status_run.font.color.rgb = RGBColor(15, 163, 107) if item.get("binding") == "confirmed" else RGBColor(199, 120, 0)
        return
    if detail_type == "next_review_plan":
        _append_docx_list_section(document, labels["outputs_until_next"], _string_list(detail.get("outputs_until_next_review")), RGBColor(17, 24, 39))
        metrics = _dict_list(detail.get("metrics_until_next_review"))
        if metrics:
            table = document.add_table(rows=1 + len(metrics), cols=3)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = labels["metric"]
            table.rows[0].cells[1].text = labels["current"]
            table.rows[0].cells[2].text = labels["next_review"]
            for row_index, row in enumerate(metrics, start=1):
                table.rows[row_index].cells[0].text = _clean_text(row.get("metric"))
                table.rows[row_index].cells[1].text = _clean_text(row.get("current"))
                table.rows[row_index].cells[2].text = _clean_text(row.get("next_review"))
            _style_docx_table(table)
        return
    for item in _string_list(detail.get("criteria")):
        document.add_paragraph(item, style="List Bullet")


def _append_docx_provenance(document: DocxDocument, provenance: NewSummaryExportProvenance) -> None:
    _append_docx_heading(document, "Provenance")
    paragraph = document.add_paragraph(_provenance_text(provenance))
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(93, 102, 117)


def _build_pdf(report: dict[str, Any], provenance: NewSummaryExportProvenance) -> bytes:
    _register_pdf_fonts()
    buffer = BytesIO()
    margin = 1.35 * cm
    frame_width = A4[0] - (2 * margin)
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=_clean_text(report["ru"].get("title") or report["en"].get("title") or "AI Summary"),
        author="Gate Challenger",
        subject="Gate Challenger AI Summary export",
    )
    styles = _pdf_styles()
    story: list[Any] = []
    for index, language in enumerate(("en", "ru")):
        if index:
            story.append(PageBreak())
        _append_pdf_version(story, report[language], language, styles, frame_width)
    _append_pdf_provenance(story, provenance, styles)
    document.build(story)
    return buffer.getvalue()


def _append_pdf_version(
    story: list[Any],
    content: dict[str, Any],
    language: str,
    styles: dict[str, ParagraphStyle],
    frame_width: float,
) -> None:
    labels = _labels(language)
    story.append(Paragraph(_xml(content.get("title") or "AI Summary"), styles["title"]))
    story.append(Paragraph(f"<b>{_xml(labels['stage'])}:</b> {_xml(content.get('stage') or 'Unknown')}", styles["body"]))
    story.append(Spacer(1, 8))
    _append_pdf_traction(story, content, labels, styles, frame_width)
    _append_pdf_text_section(story, labels["context"], [_clean_text(content.get("context"))], styles)
    _append_pdf_required(story, content, labels, styles)
    _append_pdf_list_section(story, labels["critical"], _string_list(content.get("critical_problems")), "#B91C1C", styles)
    _append_pdf_list_section(story, labels["other"], _string_list(content.get("other")), _MUTED, styles)
    _append_pdf_details(story, content, labels, styles, frame_width)


def _append_pdf_heading(story: list[Any], title: str, styles: dict[str, ParagraphStyle]) -> None:
    story.append(Spacer(1, 8))
    story.append(Paragraph(_xml(title), styles["heading"]))


def _append_pdf_text_section(story: list[Any], title: str, items: list[str], styles: dict[str, ParagraphStyle]) -> None:
    values = [item for item in items if item]
    if not values:
        return
    _append_pdf_heading(story, title, styles)
    for item in values:
        story.append(Paragraph(_xml(item), styles["body"]))


def _append_pdf_list_section(
    story: list[Any],
    title: str,
    items: list[str],
    color: str,
    styles: dict[str, ParagraphStyle],
) -> None:
    if not items:
        return
    if title:
        _append_pdf_heading(story, title, styles)
    story.append(
        ListFlowable(
            [ListItem(Paragraph(f'<font color="{color}">{_xml(item)}</font>', styles["body"])) for item in items],
            bulletType="bullet",
            leftIndent=15,
        )
    )


def _append_pdf_traction(
    story: list[Any],
    content: dict[str, Any],
    labels: dict[str, str],
    styles: dict[str, ParagraphStyle],
    frame_width: float,
) -> None:
    traction = content.get("traction_summary")
    if not _has_traction(traction):
        return
    _append_pdf_heading(story, labels["traction"], styles)
    periods = [str(item) for item in traction["periods"]]
    data = [[Paragraph(_xml(traction.get("metric_label")), styles["table_header"])] + [Paragraph(_xml(item), styles["table_header"]) for item in periods]]
    for row in traction["rows"]:
        values = row.get("values") if isinstance(row.get("values"), list) else []
        data.append(
            [Paragraph(_xml(row.get("label")), styles["table_header"])]
            + [Paragraph(_xml(values[index] if index < len(values) else ""), styles["table_body"]) for index, _ in enumerate(periods)]
        )
    table = Table(data, colWidths=_pdf_table_widths(len(data[0]), frame_width), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_SURFACE)),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(_LINE)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)


def _append_pdf_required(story: list[Any], content: dict[str, Any], labels: dict[str, str], styles: dict[str, ParagraphStyle]) -> None:
    items = content.get("required_elements")
    if not isinstance(items, list) or not items:
        return
    _append_pdf_heading(story, labels["required"], styles)
    for item in items:
        if not isinstance(item, dict):
            continue
        color = {
            "present": _SUCCESS,
            "partial": _WARNING,
            "missing": _DANGER,
            "fraction": _TEXT,
        }[_required_status_kind(item)]
        status = _localized_status(item, labels)
        story.append(
            Paragraph(
                f"<b>{_xml(item.get('label'))}</b> — <font color=\"{color}\"><b>{_xml(status)}</b></font>",
                styles["required"],
            )
        )
        story.append(Paragraph(f'<font color="{_MUTED}">{_xml(item.get("evidence"))}</font>', styles["evidence"]))


def _append_pdf_details(
    story: list[Any],
    content: dict[str, Any],
    labels: dict[str, str],
    styles: dict[str, ParagraphStyle],
    frame_width: float,
) -> None:
    entries = _ordered_details(content)
    if not entries:
        return
    _append_pdf_heading(story, labels["appendices"], styles)
    for index, (detail_id, detail) in enumerate(entries, start=1):
        story.append(Paragraph(_xml(_appendix_title(content, detail_id, index)), styles["subheading"]))
        _append_pdf_detail(story, detail, labels, styles, frame_width)


def _append_pdf_detail(
    story: list[Any],
    detail: dict[str, Any],
    labels: dict[str, str],
    styles: dict[str, ParagraphStyle],
    frame_width: float,
) -> None:
    detail_type = detail.get("type")
    if detail_type == "solution_validation":
        items = []
        for item in _dict_list(detail.get("items")):
            color = _SUCCESS if item.get("verdict") == "confirmed" else _WARNING
            items.append(Paragraph(f"{_xml(item.get('text'))} — <font color=\"{color}\"><b>{_xml(_verdict_status(item.get('verdict'), labels))}</b></font>", styles["body"]))
        if items:
            story.append(ListFlowable([ListItem(item) for item in items], bulletType="bullet", leftIndent=15))
        return
    if detail_type == "metric_binding":
        for title, key in ((labels["input_metrics"], "input_metrics"), (labels["output_metrics"], "output_metrics")):
            metrics = _dict_list(detail.get(key))
            if not metrics:
                continue
            story.append(Paragraph(_xml(title), styles["subheading"]))
            entries = []
            for item in metrics:
                color = _SUCCESS if item.get("binding") == "confirmed" else _WARNING
                entries.append(
                    Paragraph(
                        f"<b>{_xml(item.get('metric'))}</b>"
                        f"{' - ' + _xml(item.get('evidence')) if item.get('evidence') else ''}"
                        f" — <font color=\"{color}\"><b>{_xml(_detail_status(item.get('binding'), labels))}</b></font>",
                        styles["body"],
                    )
                )
            story.append(ListFlowable([ListItem(item) for item in entries], bulletType="bullet", leftIndent=15))
        return
    if detail_type == "next_review_plan":
        _append_pdf_list_section(story, labels["outputs_until_next"], _string_list(detail.get("outputs_until_next_review")), _TEXT, styles)
        metrics = _dict_list(detail.get("metrics_until_next_review"))
        if metrics:
            rows = [
                [
                    Paragraph(_xml(labels["metric"]), styles["table_header"]),
                    Paragraph(_xml(labels["current"]), styles["table_header"]),
                    Paragraph(_xml(labels["next_review"]), styles["table_header"]),
                ]
            ]
            rows.extend(
                [
                    Paragraph(_xml(row.get("metric")), styles["table_body"]),
                    Paragraph(_xml(row.get("current")), styles["table_body"]),
                    Paragraph(_xml(row.get("next_review")), styles["table_body"]),
                ]
                for row in metrics
            )
            table = Table(rows, colWidths=_pdf_table_widths(3, frame_width), repeatRows=1)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(_LINE)), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_SURFACE)), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(table)
        return
    _append_pdf_list_section(story, "", _string_list(detail.get("criteria")), _TEXT, styles)


def _append_pdf_provenance(
    story: list[Any],
    provenance: NewSummaryExportProvenance,
    styles: dict[str, ParagraphStyle],
) -> None:
    story.append(Spacer(1, 8))
    story.append(Paragraph("Provenance", styles["heading"]))
    story.append(Paragraph(_xml(_provenance_text(provenance)), styles["evidence"]))


def _labels(language: str) -> dict[str, str]:
    if language == "en":
        return {
            "appendices": "Appendices",
            "binding_confirmed": "Binding confirmed",
            "binding_insufficient": "Binding not sufficiently confirmed",
            "context": "Initiative context",
            "critical": "Identified problems",
            "current": "Current value",
            "input_metrics": "Input metrics",
            "metric": "Metric",
            "missing": "Missing",
            "partial": "Partially confirmed",
            "next_review": "Next review value",
            "other": "Other observations",
            "output_metrics": "Output metrics",
            "outputs_until_next": "Outputs until the next review",
            "present": "Present",
            "required": "Required document elements",
            "stage": "Initiative stage",
            "traction": "Traction Summary",
            "verdict_confirmed": "Confirmed",
            "verdict_insufficient": "Not sufficiently confirmed",
        }
    return {
        "appendices": "Appendices",
        "binding_confirmed": "Связь подтверждена",
        "binding_insufficient": "Связь недостаточно подтверждена",
        "context": "Краткий контекст инициативы",
        "critical": "Выявленные проблемы",
        "current": "Текущее значение",
        "input_metrics": "Input metrics",
        "metric": "Метрика",
        "missing": "Нет",
        "partial": "Частично подтверждено",
        "next_review": "Значение к следующему ревью",
        "other": "Другие наблюдения",
        "output_metrics": "Output metrics",
        "outputs_until_next": "Outputs until the next Review",
        "present": "Есть",
        "required": "Обязательные элементы документа",
        "stage": "Стадия инициативы",
        "traction": "Traction Summary",
        "verdict_confirmed": "Подтверждено",
        "verdict_insufficient": "Недостаточно подтверждено",
    }


def _has_traction(value: Any) -> bool:
    return (
        isinstance(value, dict)
        and isinstance(value.get("periods"), list)
        and bool(value["periods"])
        and isinstance(value.get("rows"), list)
        and bool(value["rows"])
    )


def _ordered_details(content: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    details = content.get("required_details")
    if not isinstance(details, dict):
        return []
    ids = [item.get("id") for item in _dict_list(content.get("required_elements")) if isinstance(item.get("id"), str)]
    ordered_ids = [*ids, *sorted(key for key in details if key not in ids)]
    return [(key, details[key]) for key in ordered_ids if isinstance(details.get(key), dict)]


def _appendix_title(content: dict[str, Any], detail_id: str, index: int) -> str:
    for item in _dict_list(content.get("required_elements")):
        if item.get("id") == detail_id and item.get("label"):
            return f"Appendix {index}. {_clean_text(item.get('label'))}"
    return f"Appendix {index}"


def _required_status_kind(item: dict[str, Any]) -> str:
    status = str(item.get("status", "")).lower()
    if status in {"есть", "present"}:
        return "present"
    if status in {"частично подтверждено", "partially confirmed"}:
        return "partial"
    numerator, separator, denominator = status.partition("/")
    if separator and numerator.isdecimal() and denominator.isdecimal():
        return "fraction"
    return "missing"


def _localized_status(item: dict[str, Any], labels: dict[str, str]) -> str:
    kind = _required_status_kind(item)
    return str(item.get("status")) if kind == "fraction" else labels[kind]


def _detail_status(status: Any, labels: dict[str, str]) -> str:
    return labels["binding_confirmed"] if status == "confirmed" else labels["binding_insufficient"]


def _verdict_status(status: Any, labels: dict[str, str]) -> str:
    return labels["verdict_confirmed"] if status == "confirmed" else labels["verdict_insufficient"]


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_clean_text(item) for item in value if _clean_text(item)]


def _dict_list(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\x00", "").split())


def _xml(value: Any) -> str:
    return escape(_clean_text(value))


def _safe_filename(title: str) -> str:
    value = "".join(character if character.isalnum() or character in " ._-" else "_" for character in title).strip()
    return value[:120] or "AI Summary"


def _provenance_text(provenance: NewSummaryExportProvenance) -> str:
    values = [
        f"analysis_id={provenance.analysis_id}",
        f"document_id={provenance.document_id}",
        f"skill={provenance.skill_name}@{provenance.skill_version}",
        f"provider_model={provenance.provider}/{provenance.model}",
    ]
    if provenance.source_revision:
        values.append(f"new_summary_source_revision={provenance.source_revision}")
    return "; ".join(values)


def _short_provenance_text(provenance: NewSummaryExportProvenance) -> str:
    return f"analysis_id={provenance.analysis_id}; document_id={provenance.document_id}"


def _pdf_table_widths(column_count: int, frame_width: float) -> list[float]:
    if column_count <= 1:
        return [frame_width]
    first = frame_width * 0.32
    rest = (frame_width - first) / (column_count - 1)
    return [first, *([rest] * (column_count - 1))]


def _register_pdf_fonts() -> None:
    if "GateSummaryRegular" in pdfmetrics.getRegisteredFontNames():
        return
    regular, bold = _find_font_paths()
    pdfmetrics.registerFont(TTFont("GateSummaryRegular", str(regular)))
    pdfmetrics.registerFont(TTFont("GateSummaryBold", str(bold or regular)))
    addMapping("GateSummary", 0, 0, "GateSummaryRegular")
    addMapping("GateSummary", 1, 0, "GateSummaryBold")
    addMapping("GateSummary", 0, 1, "GateSummaryRegular")
    addMapping("GateSummary", 1, 1, "GateSummaryBold")


def _find_font_paths() -> tuple[Path, Path | None]:
    regular_candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/local/share/fonts/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ]
    bold_candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        Path("/usr/local/share/fonts/DejaVuSans-Bold.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ]
    regular = next((path for path in regular_candidates if path.is_file()), None)
    if regular is None:
        raise NewSummaryExportUnavailableError("No Unicode font is available for PDF export")
    bold = next((path for path in bold_candidates if path.is_file()), None)
    return regular, bold


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "GateSummaryTitle",
            parent=base["Title"],
            fontName="GateSummaryBold",
            fontSize=18,
            leading=22,
            alignment=TA_LEFT,
            textColor=colors.HexColor(_TEXT),
            spaceAfter=8,
        ),
        "heading": ParagraphStyle(
            "GateSummaryHeading",
            parent=base["Heading2"],
            fontName="GateSummaryBold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor(_TEXT),
            spaceAfter=5,
        ),
        "subheading": ParagraphStyle(
            "GateSummarySubheading",
            parent=base["Heading3"],
            fontName="GateSummaryBold",
            fontSize=10,
            leading=13,
            textColor=colors.HexColor(_ACCENT),
            spaceBefore=5,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "GateSummaryBody",
            parent=base["BodyText"],
            fontName="GateSummaryRegular",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor(_TEXT),
            spaceAfter=4,
            splitLongWords=1,
        ),
        "required": ParagraphStyle(
            "GateSummaryRequired",
            parent=base["BodyText"],
            fontName="GateSummaryRegular",
            fontSize=9.5,
            leading=13,
            leftIndent=12,
            textColor=colors.HexColor(_TEXT),
            spaceAfter=1,
            splitLongWords=1,
        ),
        "evidence": ParagraphStyle(
            "GateSummaryEvidence",
            parent=base["BodyText"],
            fontName="GateSummaryRegular",
            fontSize=8,
            leading=10,
            leftIndent=12,
            textColor=colors.HexColor(_MUTED),
            spaceAfter=5,
            splitLongWords=1,
        ),
        "table_header": ParagraphStyle(
            "GateSummaryTableHeader",
            parent=base["BodyText"],
            fontName="GateSummaryBold",
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor(_TEXT),
            splitLongWords=1,
        ),
        "table_body": ParagraphStyle(
            "GateSummaryTableBody",
            parent=base["BodyText"],
            fontName="GateSummaryRegular",
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor(_TEXT),
            splitLongWords=1,
        ),
    }
