from copy import deepcopy
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.services.new_summary_exports import NewSummaryExportProvenance, _build_docx, _build_pdf
from app.services.new_summary_traction import with_traction_totals, traction_row_label


def _table(metric, periods, values, **extra):
    return {"metric": metric, "metric_label": f"{metric.upper()} (incr)", "periods": periods,
            "rows": [{"label": "Total incremental output uplifts", "values": values}], **extra}


def _align(revenue, dtb):
    payload = {"language": "en", "traction_summary": {"tables": [revenue, dtb]}}
    before = deepcopy(payload)
    result = with_traction_totals(payload)
    assert payload == before
    assert with_traction_totals(result) == result
    return result["traction_summary"]["tables"]


def test_later_start_annual_granularity_and_independent_tail():
    revenue, dtb = _align(
        _table("revenue", ["2024", "2025", "Q3 2026", "Q4 2026", "2026", "2027", "2028", "Total"],
               ["1", "2", "3", "4", "7", "8", "9", "27"]),
        _table("dtb", ["CY26", "CY27", "CY28", "CY29", "Total"], ["1%", "2%", "3%", "4%", "8%"]),
    )
    assert revenue["periods"] == ["2026", "2027", "2028", "Total 2026–2028"]
    assert revenue["rows"][0]["values"] == ["7", "8", "9", "could not extract data"]
    assert dtb["periods"] == ["2026", "2027", "2028", "2029", "Total"]
    assert dtb["rows"][0]["values"] == ["1%", "2%", "3%", "4%", "8%"]


def test_p2p_h2_fallback_is_explicit_and_source_total_is_kept():
    revenue, dtb = _align(
        _table("revenue", ["2024", "2025", "2026", "2027", "2028", "2029", "2030", "2026-30 total"],
               ["-", "-", "-", "202", "1372", "3153", "5157", "9884"]),
        _table("dtb", ["2026 H2", "2027", "2028", "2029", "2030"], ["0%", "0.47%", "1.66%", "2.94%", "3.91%"], cumulative=True),
    )
    assert revenue["periods"][:-1] == dtb["periods"][:-1] == ["2026", "2027", "2028", "2029", "2030"]
    assert revenue["rows"][0]["values"][-1] == "9884"
    assert dtb["rows"][0]["values"][0] == "0% (H2 2026)"
    assert dtb["rows"][0]["values"][-1] == "—"


def test_full_year_takes_priority_over_h2_and_total_is_not_recalculated():
    revenue, dtb = _align(
        _table("revenue", ["2026", "2027", "Total"], ["10", "20", "30"]),
        _table("dtb", ["H2 2026", "2026", "2027", "Total"], ["0.5%", "0.25%", "1.25%", "1.5%"]),
    )
    assert dtb["periods"] == revenue["periods"]
    assert dtb["rows"][0]["values"] == ["0.25%", "1.25%", "1.5%"]


def test_unknown_period_or_conflicting_aliases_are_not_silently_removed():
    for periods in (["Jan 2026", "2027"], ["2026", "CY26"]):
        revenue = _table("revenue", periods, ["10", "20"])
        result, _ = _align(revenue, _table("dtb", ["2026", "2027"], ["1%", "2%"]))
        assert result["periods"][:-1] == periods
        assert result["rows"][0]["values"][:2] == ["10", "20"]


def test_metric_label_is_in_body_and_custom_row_labels_are_preserved():
    table = _table("revenue", ["2026"], ["10"])
    assert traction_row_label(table, table["rows"][0]) == "REVENUE (incr)"
    assert traction_row_label(table, {"label": "Scenario A"}) == "REVENUE (incr): Scenario A"


def test_v4_exports_render_headers_plan_fact_and_safe_links():
    content = {"schema_version": "new-summary-v4", "language": "en", "title": "AI Summary Example",
               "stage": "Progress Review", "context": "Example context.", "critical_problems": [],
               "traction_summary": {"tables": [_table("revenue", ["2026", "Total"], ["10", "10"])]},
               "required_elements": [
                   {"id": "progress_review_plan_fact_last_half_year", "label": "Plan vs actual", "status": "есть",
                    "detail": {"type": "plan_fact", "launches": [{"output": "Pilot", "status": "partial", "comment": "One market."}],
                               "metrics": [{"metric": "Users", "planned": "100", "actual": "80"}]}},
                   {"id": "gate2_user_flow", "label": "Mockups", "status": "есть", "detail": {
                       "type": "source_links", "availability": "provided", "links": [
                           {"label": "Prototype", "url": "https://example.com/mockup"},
                           {"label": "Unsafe", "url": "javascript:alert(1)"},
                       ]}},
               ]}
    report = {"ru": {**content, "language": "ru"}, "en": content}
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    docx = Document(BytesIO(_build_docx(report, provenance)))
    assert docx.tables[0].cell(0, 0).text == "Output metrics — uplifts"
    assert docx.tables[0].cell(1, 0).text == "REVENUE (incr)"
    assert [cell.text for cell in docx.tables[1].rows[0].cells] == ["Metric", "Plan", "Actual"]
    assert any("Pilot — Partially completed" in p.text for p in docx.paragraphs)
    links = [rel.target_ref for rel in docx.part.rels.values() if rel.is_external]
    assert "https://example.com/mockup" in links and not any("javascript:" in url for url in links)
    pdf = PdfReader(BytesIO(_build_pdf(report, provenance)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Output metrics" in text and "Launches: plan vs actual" in text and "Prototype" in text
    assert "javascript" not in text
