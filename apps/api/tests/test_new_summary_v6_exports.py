from copy import deepcopy
from io import BytesIO

from docx import Document
from pypdf import PdfReader
import pytest

from app.services.new_summary_exports import NewSummaryExportProvenance, _build_docx, _build_pdf


@pytest.mark.parametrize("language", ["ru", "en"])
def test_v6_exports_counts_stops_and_launch_colors(language):
    content = {"schema_version": "new-summary-v6", "language": language, "title": "AI Summary Example",
        "stage": "Progress Review", "context": "Context.", "traction_summary": {"tables": []}, "required_elements": [
            {"id": "gate2_metric_linkage", "label": "Metrics", "status": "частично подтверждено", "detail": {"type": "metric_binding",
                "input_metrics": [{"metric": "Conversion", "binding": "confirmed", "evidence": "Target outcome."}],
                "output_metrics": [{"metric": "Volume", "binding": "insufficient", "evidence": "No linkage."}]}},
            {"id": "gate2_stop_criteria", "label": "Stop criteria", "status": "есть", "evidence": "Stop if pilot fails.",
                "detail": {"type": "stop_criteria", "primary_criterion": "Stop if pilot fails.", "criteria": ["Stop if funding ends."]}},
            {"id": "gate3_stop_criteria", "label": "Stop criteria", "status": "есть", "evidence": "Stop if pilot fails.",
                "detail": {"type": "stop_criteria", "criteria": []}},
            {"id": "progress_review_plan_fact_last_half_year", "label": "Plan vs actual", "status": "частично подтверждено",
                "detail": {"type": "plan_fact", "metrics": [], "launches": [
                    {"output": "Alpha", "status": "completed", "comment": "Delivered."},
                    {"output": "Beta", "status": "partial"}, {"output": "Gamma", "status": "not_completed"},
                    {"output": "Delta", "status": "unknown"}]}},
        ]}
    report = {"en": content, "ru": content}
    original = deepcopy(report)
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    docx = Document(BytesIO(_build_docx(report, provenance)))
    count = "связь 1 метрики из 2 подтверждена" if language == "ru" else "binding of 1 of 2 metrics confirmed"
    heading = next(p for p in docx.paragraphs if "Metrics:" in p.text and count in p.text)
    assert heading.runs[0].bold and heading.runs[1].bold is False
    assert count in heading.text
    additional = "Дополнительные найденные Stop критерии" if language == "ru" else "Additional stop criteria"
    assert sum(p.text == additional for p in docx.paragraphs) == 1
    for output, color in (("Alpha", "0FA36B"), ("Beta", "C77800"), ("Gamma", "D92D20"), ("Delta", "5D6675")):
        paragraph = next(p for p in docx.paragraphs if p.text.startswith(output))
        assert str(paragraph.runs[1].font.color.rgb) == color
    pdf = PdfReader(BytesIO(_build_pdf(report, provenance)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert count in text.replace("\n", " ")
    assert text.count(additional) == 1
    assert "Stop if funding ends." in text
    assert report == original
