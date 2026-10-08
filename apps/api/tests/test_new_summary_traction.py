from copy import deepcopy
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.services.new_summary_traction import needs_verified_revenue_total, with_traction_totals


def test_v7_single_table_keeps_four_rows_and_mismatch_without_legacy_rewrite():
    table = {"periods": ["2026", "Total"], "rows": [
        {"label": "DTB Uplift (Cumul)", "values": ["2%", "—"]},
        {"label": "Revenue from DTB", "values": ["10", "10"]},
        {"label": "Revenue non-DTB", "values": ["20", "20"]},
        {"label": "Total Revenue", "values": ["35", "35"], "mismatch_periods": ["2026", "Total"]},
    ]}
    payload = {"schema_version": "new-summary-v7", "language": "ru", "traction_summary": table}

    result = with_traction_totals(payload)

    assert result["traction_summary"] == table
    assert result is not payload


def test_historical_traction_break_marker_is_cleaned_without_changing_total():
    payload = {"language": "ru", "traction_summary": {"tables": [{
        "metric": "revenue", "metric_label": "Выручка (инкр.)",
        "periods": ["2026", "2026-31<br>total"],
        "rows": [{"label": "Revenue uplift", "values": ["10", "15 918"]}],
    }]}}
    result = with_traction_totals(payload)["traction_summary"]["tables"][0]
    assert result["periods"] == ["2026", "2026-31 total"]
    assert result["rows"][0]["values"] == ["10", "15 918"]
    assert payload["traction_summary"]["tables"][0]["periods"][-1] == "2026-31<br>total"
from app.services.new_summary_exports import (
    NewSummaryExportProvenance, _build_docx, _build_pdf, _confirmed_first,
    _gate2_hypothesis_heading, _labels, _validation_rationale,
)
from app.services.new_summaries import with_verified_source_totals
from app.schemas.analyses import NewSummaryRead, NewSummaryVariantRead
from uuid import uuid4


def _payload(language: str = "ru") -> dict:
    return {
        "language": language,
        "traction_summary": {"tables": [
            {"metric": "revenue", "metric_label": "Revenue", "periods": ["2026", "2027"],
             "rows": [{"label": "Incremental revenue", "values": ["10", "20"]}]},
            {"metric": "dtb", "metric_label": "DTB", "periods": ["CY26", "CY27"],
             "rows": [{"label": "DTB uplift", "values": ["1%", "2%"]}]},
        ]},
    }


def test_missing_total_is_explicit_and_never_calculated():
    payload = _payload()
    original = deepcopy(payload)
    tables = with_traction_totals(payload)["traction_summary"]["tables"]
    assert tables[0]["periods"] == ["2026", "2027", "Total 2026–2027"]
    assert tables[0]["rows"][0]["values"] == ["10", "20", "невозможно извлечь данные"]
    assert tables[1]["periods"][-1] == "Total"
    assert payload == original


def test_verified_source_without_total_uses_document_absence_reason():
    payload = _payload()
    source = [{"metric": "revenue", "periods": ["2026", "2027"]}]
    tables = with_traction_totals(payload, source_tables=source)["traction_summary"]["tables"]
    assert tables[0]["rows"][0]["values"][-1] == "отсутствуют данные в документе защиты"
    assert tables[1]["rows"][0]["values"][-1] == "невозможно извлечь данные"


def test_existing_total_values_are_preserved_in_both_languages():
    for language in ("ru", "en"):
        payload = _payload(language)
        table = payload["traction_summary"]["tables"][0]
        table["periods"].append("2026–2027 total")
        table["rows"][0]["values"].append("17")
        result = with_traction_totals(payload)["traction_summary"]["tables"]
        assert result[0] == table
        assert len(result[0]["periods"]) == len(result[0]["rows"][0]["values"])
        assert result[1]["rows"][0]["values"][-1] == (
            "невозможно извлечь данные" if language == "ru" else "could not extract data"
        )


def test_existing_total_moves_to_last_column_with_matching_value():
    payload = _payload()
    table = payload["traction_summary"]["tables"][0]
    table["periods"] = ["2026", "Total", "2027"]
    table["rows"][0]["values"] = ["10", "30", "20"]
    result = with_traction_totals(payload)["traction_summary"]["tables"][0]
    assert result["periods"] == ["2026", "2027", "Total"]
    assert result["rows"][0]["values"] == ["10", "20", "30"]


def test_saved_summary_placeholder_is_enriched_only_by_verified_source_total():
    payload = _payload()
    table = payload["traction_summary"]["tables"][0]
    table["periods"].append("Total 2026–2027")
    table["rows"][0]["values"].append("невозможно извлечь данные")
    source = [{"id": "faq6", "metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"],
        ["Revenue", "10", "20", "30"],
    ]}}]
    enriched = with_traction_totals(payload, source_blocks=source)
    assert enriched["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "30"
    assert payload["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "невозможно извлечь данные"
    source[0]["metadata"]["rows"][1][2] = "21"
    assert with_traction_totals(payload, source_blocks=source)["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "невозможно извлечь данные"


def test_old_bilingual_report_reads_verified_total_without_mutating_saved_payload(monkeypatch):
    import app.services.new_summaries as summaries

    source = [{"id": "faq6", "metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"],
        ["Revenue", "10", "20", "30"],
    ]}}]
    monkeypatch.setattr(summaries, "verified_table_blocks", lambda _document: source)
    ru = _payload("ru")
    en = _payload("en")
    report = NewSummaryRead(
        analysis_id=uuid4(), ru=NewSummaryVariantRead(status="completed", payload=ru),
        en=NewSummaryVariantRead(status="completed", payload=en),
    )
    updated = with_verified_source_totals(report, object())
    for language in ("ru", "en"):
        assert getattr(updated, language).payload["traction_summary"]["tables"][0]["rows"][0]["values"][-1] == "30"
        assert getattr(report, language).payload["traction_summary"]["tables"][0]["periods"] == ["2026", "2027"]


def test_gate2_hypothesis_heading_uses_actual_items_and_confirmed_first():
    item = {"id": "gate2_hypothesis_results", "detail": {"type": "solution_validation", "items": [
        {"text": "a", "verdict": "insufficient"}, {"text": "b", "verdict": "confirmed"},
        {"text": "c", "verdict": "insufficient"},
    ]}}
    assert _gate2_hypothesis_heading(item, _labels("ru")) == (
        "Результаты проверки гипотез из Gate: 1 гипотез из 3 подтверждены, 2 гипотез из 3 недостаточно подтверждены."
    )
    assert [entry["text"] for entry in _confirmed_first(item["detail"]["items"], "verdict")] == ["b", "a", "c"]


def test_v2_solution_checks_show_counts_and_source_rationale():
    item = {"id": "stream_review_1_solution_validation", "detail": {"type": "solution_validation", "items": [
        {"text": "Pilot", "verdict": "confirmed", "test": "Pilot with 20 users",
         "expected_result": "10 activations", "actual_result": "12 activations"},
        {"text": "Launch", "verdict": "insufficient"},
    ]}}
    assert _gate2_hypothesis_heading(item, _labels("ru"), new_format=True) == (
        "Подтвержденные решения: 1 проверка из 2 подтверждена, 1 проверка из 2 недостаточно подтверждена."
    )
    assert _validation_rationale(item["detail"]["items"][0], _labels("ru")) == (
        "Проверка: Pilot with 20 users; ожидали: 10 activations; получили: 12 activations."
    )


def test_v2_docx_omits_removed_sections_and_numbers_required_elements():
    item = {"id": "stream_review_1_solution_validation", "label": "Подтверждение решения", "status": "1/1",
            "detail": {"type": "solution_validation", "items": [{
                "text": "Пилот", "verdict": "confirmed", "test": "Пилот на 20 пользователях",
                "expected_result": "10 активаций", "actual_result": "12 активаций",
            }]}}
    content = {"schema_version": "new-summary-v2", "language": "ru", "title": "AI Summary Test",
               "stage": "Stream Review 1", "context": "Проверен пилот.",
               "traction_summary": {"tables": []}, "required_elements": [item],
               "document_quality_percent": 75, "other": ["Legacy-only note"], "critical_problems": []}
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    document = DocxDocument(BytesIO(_build_docx({"ru": content, "en": {**content, "language": "en"}}, provenance)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "1. Подтвержденные решения: 1 проверка из 1 подтверждена" in text
    assert "Проверка: Пилот на 20 пользователях" in text
    assert "Качество документа" not in text
    assert "Другие наблюдения" not in text
    assert "Legacy-only note" not in text


def test_v3_problem_issue_is_bold_in_exports_and_v2_string_still_renders():
    content = {
        "schema_version": "new-summary-v3", "language": "en", "title": "AI Summary Test",
        "stage": "Gate 2", "context": "A pilot was conducted.",
        "traction_summary": {"tables": []}, "required_elements": [],
        "critical_problems": [{
            "issue": "Revenue starts late.",
            "fact": "The current plan shows no revenue in the first year.",
        }],
    }
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    report = {"ru": {**content, "language": "ru"}, "en": content}
    document = DocxDocument(BytesIO(_build_docx(report, provenance)))
    problem = next(paragraph for paragraph in document.paragraphs if "Revenue starts late" in paragraph.text)
    assert problem.runs[0].text == "Revenue starts late."
    assert problem.runs[0].bold is True
    assert problem.runs[1].text == " The current plan shows no revenue in the first year."
    assert not problem.runs[1].bold
    assert "Revenue starts late." in PdfReader(BytesIO(_build_pdf(report, provenance))).pages[0].extract_text()

    legacy = {**content, "schema_version": "new-summary-v2", "critical_problems": ["Legacy plain-text problem."]}
    legacy_report = {"ru": {**legacy, "language": "ru"}, "en": legacy}
    legacy_document = DocxDocument(BytesIO(_build_docx(legacy_report, provenance)))
    assert any("Legacy plain-text problem." in paragraph.text for paragraph in legacy_document.paragraphs)


def test_legacy_single_revenue_table_without_metric_can_receive_verified_total():
    payload = {"language": "ru", "traction_summary": {
        "metric_label": "Revenue (incr)", "periods": ["2026", "2027", "Total"],
        "rows": [{"label": "Revenue", "values": ["10", "20", "Не смог получить данные"]}],
    }}
    source = [{"metadata": {"rows": [
        ["ToBe P&L", "2026", "2027", "2026-27 total"], ["Revenue", "10", "20", "30"],
    ]}}]
    assert needs_verified_revenue_total(payload)
    assert with_traction_totals(payload, source_blocks=source)["traction_summary"]["rows"][0]["values"][-1] == "30"
