from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.services.new_summary_exports import NewSummaryExportProvenance, _build_docx, _build_pdf


def test_v7_exports_one_table_with_four_rows_and_mismatch_note():
    table = {"periods": ["2026", "Total"], "rows": [
        {"label": "DTB Uplift (Cumul)", "values": ["2%", "—"]},
        {"label": "Revenue from DTB", "values": ["10", "10"]},
        {"label": "Revenue non-DTB", "values": ["20", "20"]},
        {"label": "Total Revenue", "values": ["35", "35"], "mismatch_periods": ["2026"]},
    ]}
    report = {language: {
        "schema_version": "new-summary-v7", "language": language,
        "title": "AI Summary Example", "stage": "Gate 2", "context": "Context.",
        "traction_summary": table, "required_elements": [],
    } for language in ("ru", "en")}
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)

    document = Document(BytesIO(_build_docx(report, provenance)))
    for table in document.tables:
        if table.cell(0, 0).text.startswith("Output") and "Не равно" in table.cell(4, 1).text:
            assert len(table.rows) == 5
            assert table.cell(4, 0).text == "Total Revenue"
            assert "Не равно Rev. from DTB + Rev. non-DTB" in table.cell(4, 1).text
            break
    else:
        raise AssertionError("traction table missing")

    pdf_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(_build_pdf(report, provenance))).pages)
    assert "DTB Uplift (Cumul)" in pdf_text
    assert "Revenue non-DTB" in pdf_text
    assert "Не равно Rev. from DTB + Rev. non-DTB" in pdf_text
