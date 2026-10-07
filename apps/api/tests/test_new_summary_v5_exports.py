from io import BytesIO

from docx import Document
from pypdf import PdfReader
import pytest

from app.services.new_summary_exports import NewSummaryExportProvenance, _build_docx, _build_pdf


@pytest.mark.parametrize("scheme", ["https", "HTTPS"])
def test_v5_exports_keep_mixed_weight_headings_and_inline_metric_evidence(scheme):
    url = f"{scheme}://example.com/prototype?node=1&mode=design"
    content = {"schema_version": "new-summary-v5", "language": "en", "title": "AI Summary Example",
        "stage": "Gate 2", "context": "Context.", "traction_summary": {"tables": []}, "required_elements": [
            {"id": "gate2_hypothesis_results", "label": "Hypotheses", "status": "1/1", "detail": {
                "type": "solution_validation", "items": [{"text": "The pilot validated demand.", "verdict": "confirmed"}]}},
            {"id": "stream_review_1_solution_validation", "label": "Solutions", "status": "1/1", "detail": {
                "type": "solution_validation", "items": [{"text": "The prototype validated the solution.", "verdict": "confirmed"}]}},
            {"id": "gate2_metric_linkage", "label": "Metrics", "status": "есть", "detail": {"type": "metric_binding",
                "input_metrics": [{"metric": "Conversion", "binding": "confirmed", "evidence": "Measures the target outcome."}], "output_metrics": []}},
            {"id": "gate2_user_flow", "label": "Mockups", "status": "есть", "detail": {"type": "source_links", "availability": "provided",
                "links": [{"label": "Design", "url": url}]}},
        ]}
    report = {"en": content, "ru": {**content, "language": "ru"}}
    provenance = NewSummaryExportProvenance("a", "d", "skill", "1", "test", "mock", None)
    docx = Document(BytesIO(_build_docx(report, provenance)))
    heading = next(p for p in docx.paragraphs if "Gate 1 hypothesis validation:" in p.text)
    assert heading.runs[0].bold and heading.runs[0].text.endswith(":")
    assert heading.runs[1].text.startswith(" 1 of 1") and heading.runs[1].bold is False
    metric = next(p for p in docx.paragraphs if p.text.startswith("Conversion"))
    assert metric.runs[0].bold
    assert metric.runs[1].bold is False
    assert metric.text.index("Measures") < metric.text.index("Binding is relevant")
    assert "\n" not in metric.text
    assert any(rel.target_ref == url for rel in docx.part.rels.values() if rel.is_external)
    pdf = PdfReader(BytesIO(_build_pdf(report, provenance)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "hypothesis validation: 1 of 1" in text
    assert "quantitative research, prototypes or fake doors" in text
    assert text.index("Measures the target outcome.") < text.index("Binding is relevant")
    urls = [annot.get_object().get("/A", {}).get("/URI") for page in pdf.pages for annot in page.get("/Annots", [])]
    assert url in urls
